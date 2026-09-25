import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Laboratorio\Documents\ChatGPT\9524")
SOURCE = ROOT / "propuesta_intercambio_aprendizajes.md"
OUTPUT = ROOT / "propuesta_intercambio_aprendizajes.docx"


def set_font(style, name, size, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:ascii"), name)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), name)


def keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:keepNext"))


def add_inline(paragraph, text):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
set_font(styles["Normal"], "Aptos", 11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08

set_font(styles["Title"], "Aptos Display", 20, True)
styles["Title"].paragraph_format.space_after = Pt(14)
styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

for name, size, before, after in [
    ("Heading 1", 16, 14, 6),
    ("Heading 2", 13, 12, 5),
    ("Heading 3", 11.5, 10, 4),
]:
    set_font(styles[name], "Aptos Display", size, True)
    styles[name].paragraph_format.space_before = Pt(before)
    styles[name].paragraph_format.space_after = Pt(after)
    styles[name].paragraph_format.keep_with_next = True

set_font(styles["List Bullet"], "Aptos", 11)
styles["List Bullet"].paragraph_format.left_indent = Inches(0.25)
styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.18)
styles["List Bullet"].paragraph_format.space_after = Pt(3)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
title_used = False
pending_page_break = False

for raw in lines:
    line = raw.rstrip()
    if not line:
        continue
    if line.strip() == "---":
        pending_page_break = True
        continue

    heading = re.match(r"^(#{1,6})\s+(.*)$", line)
    if heading:
        level = len(heading.group(1))
        text = heading.group(2).strip()
        if pending_page_break:
            doc.add_page_break()
            pending_page_break = False
        if not title_used and level == 1:
            p = doc.add_paragraph(style="Title")
            add_inline(p, text)
            title_used = True
        else:
            style_level = min(max(level - 1, 1), 3)
            p = doc.add_paragraph(style=f"Heading {style_level}")
            add_inline(p, text)
            keep_with_next(p)
        continue

    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, line[2:].strip())
        continue

    p = doc.add_paragraph()
    add_inline(p, line)

doc.core_properties.title = "Plataforma de intercambio de aprendizajes"
doc.core_properties.subject = "Propuesta y alcance del MVP"
doc.core_properties.author = ""
doc.save(OUTPUT)
print(OUTPUT)
